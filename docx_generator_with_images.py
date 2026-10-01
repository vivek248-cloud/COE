#!/usr/bin/env python3
"""Image-aware wrapper for the official Holy Cross question-paper DOCX generator.

The existing generator remains the layout source of truth. This wrapper only
replaces the question-row renderer so embedded diagrams/images survive into the
printable DOCX without changing the institutional 4-section layout.
"""

import base64
import io
import os
import sys

import docx_generator as base
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _image_stream(value):
    """Return a readable binary stream for a safe local/data image source."""
    if not value or not isinstance(value, str):
        return None
    value = value.strip()
    if not value:
        return None

    # DOCX importer stores embedded images as data URIs.
    if value.startswith('data:image/') and ';base64,' in value:
        try:
            encoded = value.split(';base64,', 1)[1]
            return io.BytesIO(base64.b64decode(encoded, validate=True))
        except Exception:
            return None

    if value.startswith('file://'):
        value = value[7:]

    # Resolve common project-relative upload paths. Do not fetch remote URLs.
    candidates = [value]
    if not os.path.isabs(value):
        project_dir = os.path.dirname(os.path.abspath(base.__file__))
        candidates.extend([
            os.path.join(project_dir, value),
            os.path.join(os.path.dirname(project_dir), value),
        ])

    for candidate in candidates:
        try:
            if os.path.isfile(candidate):
                with open(candidate, 'rb') as fh:
                    return io.BytesIO(fh.read())
        except Exception:
            continue
    return None


def _question_images(q):
    values = []
    for key in ('image_url', 'image_path', 'image_reference'):
        value = q.get(key)
        if isinstance(value, str) and value.strip():
            values.append(value.strip())
    images = q.get('images')
    if isinstance(images, list):
        for item in images:
            if isinstance(item, str) and item.strip():
                values.append(item.strip())
            elif isinstance(item, dict):
                for key in ('image_url', 'image_path', 'src', 'data'):
                    if isinstance(item.get(key), str) and item[key].strip():
                        values.append(item[key].strip())
                        break

    seen = set()
    out = []
    for value in values:
        marker = value[:200]
        if marker not in seen:
            seen.add(marker)
            out.append(value)
    return out


def render_questions_table_with_images(doc, questions):
    if not questions:
        return

    t_q = doc.add_table(rows=0, cols=4)
    t_q.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_q.autofit = False
    col_w = [Inches(0.6), Inches(4.8), Inches(0.8), Inches(0.8)]

    hdr_cells = t_q.add_row().cells
    hdr_titles = ["Q.No", "Question Description", "K-Level", "CO Level"]
    for idx, title in enumerate(hdr_titles):
        hdr_cells[idx].width = col_w[idx]
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = "Times New Roman"
        base.set_cell_borders(hdr_cells[idx], top="single", bottom="single", left="none", right="none")
        base.set_cell_margins(hdr_cells[idx], top=30, bottom=30, left=50, right=50)

    for q in questions:
        row_cells = t_q.add_row().cells
        for idx, w in enumerate(col_w):
            row_cells[idx].width = w
            base.set_cell_borders(row_cells[idx], top="none", bottom="none", left="none", right="none")
            base.set_cell_margins(row_cells[idx], top=30, bottom=30, left=50, right=50)

        q_num = base.clean_xml_string(str(q.get('q_number', '')))
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(q_num)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.name = "Times New Roman"

        q_text = base.clean_xml_string(q.get('question_text', ''))
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(q_text)
        r1.font.size = Pt(9.5)
        r1.font.name = "Times New Roman"

        # Render embedded diagrams/images below the question text.
        for image_source in _question_images(q):
            stream = _image_stream(image_source)
            if stream is None:
                continue
            try:
                p_img = row_cells[1].add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.add_run().add_picture(stream, width=Inches(3.8))
            except Exception:
                # A malformed image must never abort the complete paper export.
                continue

        k_val = base.clean_xml_string(q.get('k_level', 'K1'))
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(f"[{k_val}]")
        r2.bold = True
        r2.font.size = Pt(9)
        r2.font.name = "Times New Roman"

        co_val = base.clean_xml_string(q.get('co_level', 'CO1'))
        p3 = row_cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(2)
        r3 = p3.add_run(f"[{co_val}]")
        r3.bold = True
        r3.font.size = Pt(9)
        r3.font.name = "Times New Roman"


# Patch only the renderer used by base.generate_docx(). All official header,
# section, pagination and typography logic remains in docx_generator.py.
base.render_questions_table = render_questions_table_with_images


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python3 docx_generator_with_images.py <input_json_path> <output_docx_path>")
        sys.exit(1)
    try:
        import json
        with open(sys.argv[1], 'r', encoding='utf-8') as fh:
            data = json.load(fh)
        base.generate_docx(data, sys.argv[2])
        print("DOCX successfully generated:", sys.argv[2])
    except Exception as exc:
        print("Error generating DOCX:", str(exc), file=sys.stderr)
        sys.exit(1)
